"""资产库：列表 / 筛选 / CSV 导出 / Excel 导入 / 增删改 / 批量删除。"""

import csv
from urllib.parse import urlencode

from openpyxl import Workbook, load_workbook

from django.contrib import messages
from django.core.paginator import Paginator
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse

from .. import oplog
from ..models import Asset, AssetAttachment, Department
from ..permissions import perm_required
from .common import (
    ASSET_FIELD_LABELS, _asset_label, _asset_snapshot,
    _filtered_assets, _missing_required_asset, _parse_date, _parse_price,
)

_HEADER_ALIASES = {
    'asset_id': ['资产编号', '编号', '资产ID'],
    'name': ['资产名称', '名称'],
    'category': ['类别', '资产类别'],
    'specification': ['规格', '型号'],
    'configuration': ['配置', '配置参数'],
    'department': ['所属部门', '部门', '责任部门'],
    'user': ['实际使用人', '使用人'],
    'responsible': ['责任人', '责任人员'],
    'serial': ['序列号', '机身序列号', '序列号(SN)'],
    'status': ['状态', '使用状态', '资产状态'],
    'price': ['价格', '价值', '资产价值', '价格(元)'],
    'purchase_date': ['购置日期', '日期', '购置时间'],
    'location': ['位置', '存放位置', '存放地点'],
}


@perm_required('view_assets')
def library(request):
    qs, filters = _filtered_assets(request)
    status = filters['status']
    category = filters['category']
    dept_id = filters['department']
    keyword = filters['q']
    field = filters['field']
    sort = request.GET.get('sort', 'asset_id')
    direction = request.GET.get('dir', 'asc')

    # 排序白名单
    sort_map = {
        'asset_id': 'asset_id', 'name': 'name', 'category': 'category',
        'price': 'price', 'status': 'status', 'department': 'department__name',
    }
    order_field = sort_map.get(sort, 'asset_id')
    order = ('-' if direction == 'desc' else '') + order_field
    qs = qs.order_by(order)

    # 导出 CSV
    if request.GET.get('export') == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response.write('\ufeff')  # BOM，方便 Excel 识别中文
        w = csv.writer(response)
        w.writerow(['资产编号', '资产名称', '资产类别', '所属部门', '责任人', '实际使用人', '存放位置', '使用状态', '序列号', '规格', '配置', '价格(元)'])
        for a in qs:
            w.writerow([
                a.asset_id, a.name, a.category, a.department.name if a.department else '',
                a.responsible, a.user, a.location, a.status, a.serial,
                a.specification, a.configuration, a.price,
            ])
        filename = 'assets_export.csv'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    # 每页条数
    try:
        per_page = int(request.GET.get('per_page', 12))
    except (TypeError, ValueError):
        per_page = 12
    if per_page not in (10, 20, 50, 100):
        per_page = 12

    paginator = Paginator(qs, per_page)
    page_obj = paginator.get_page(request.GET.get('page'))

    # 保留筛选参数的查询串（排除排序/方向/页码），供排序链接复用
    keep = {k: v for k, v in request.GET.items() if k not in ('sort', 'dir', 'page')}
    base_qs = urlencode(keep)
    qs_no_perpage = urlencode({k: v for k, v in request.GET.items() if k not in ('page', 'per_page')})

    context = {
        'page_obj': page_obj,
        'assets': page_obj.object_list,
        'departments': Department.objects.all(),
        'category_options': [c[0] for c in Asset.CATEGORY_CHOICES],
        'status_options': [s[0] for s in Asset.STATUS_CHOICES],
        'page_title': '资产库',
        'active': 'library',
        'sort': sort,
        'direction': direction,
        'base_qs': base_qs,
        'qs_no_perpage': qs_no_perpage,
        'per_page': per_page,
        'per_page_choices': (10, 20, 50, 100),
        'filters': filters,
    }
    return render(request, 'library.html', context)


@perm_required('manage_assets')
def assets_bulk_delete(request):
    if request.method == 'POST':
        idstr = request.POST.get('ids', '').strip()
        ids = [int(x) for x in idstr.split(',') if x.strip().isdigit()]
        if ids:
            # 级联保护：跳过关联有领用/变更历史的资产，避免静默丢失审计记录
            protected = []
            deletable = []
            for a in Asset.objects.filter(pk__in=ids).select_related('department'):
                if a.requisitions.exists() or a.changes.exists():
                    protected.append(a.asset_id)
                else:
                    deletable.append(a.pk)
            deleted_n = len(deletable)
            if deleted_n:
                Asset.objects.filter(pk__in=deletable).delete()
            if deleted_n or protected:
                detail_parts = [f'已删除 {deleted_n} 项']
                if protected:
                    detail_parts.append(
                        f'因存在领用/变更历史被跳过 {len(protected)} 项：'
                        f'{", ".join(protected[:20])}'
                    )
                oplog.log(request, 'asset', 'delete',
                          target=f'批量删除（{deleted_n} 项）',
                          detail='；'.join(detail_parts))
            if protected:
                messages.warning(
                    request,
                    f'已跳过 {len(protected)} 项存在关联领用/变更历史的资产'
                    f'（{", ".join(protected[:8])}{"…" if len(protected) > 8 else ""}），请逐项处理后删除。',
                )
            if deleted_n:
                messages.success(request, f'已删除 {deleted_n} 条资产。')
            if not deleted_n and not protected:
                messages.error(request, '没有可删除的资产。')
        else:
            messages.error(request, '请先勾选要删除的资产。')
        nxt = request.POST.get('next') or reverse('library')
        return redirect(nxt)
    messages.error(request, '非法请求。')
    return redirect(reverse('library'))


@perm_required('manage_assets')
def asset_create(request):
    if request.method == 'POST':
        asset_id = request.POST.get('asset_id', '').strip()
        name = request.POST.get('name', '').strip()
        category = request.POST.get('category', '')
        specification = request.POST.get('specification', '').strip()
        configuration = request.POST.get('configuration', '').strip()
        responsible = request.POST.get('responsible', '').strip()
        serial = request.POST.get('serial', '').strip()
        department_id = request.POST.get('department') or None
        user = request.POST.get('user', '').strip()
        status = request.POST.get('status', '库存')
        price_value = _parse_price(request.POST.get('price'))
        price = price_value if price_value is not None else 0
        raw_date = (request.POST.get('purchase_date') or '').strip()
        purchase_date = _parse_date(raw_date) if raw_date else None
        location = request.POST.get('location', '').strip()

        missing = _missing_required_asset(dict(
            asset_id=asset_id, name=name, category=category, department=department_id,
            responsible=responsible, user=user, location=location, status=status,
        ))
        if missing:
            messages.error(request, '请填写必填项：' + '、'.join(missing))
        elif Asset.objects.filter(asset_id=asset_id).exists():
            messages.error(request, f'资产编号 {asset_id} 已存在。')
        elif price_value is None:
            messages.error(request, '资产价值必须是数字。')
        else:
            created = Asset.objects.create(
                asset_id=asset_id, name=name, category=category,
                specification=specification, configuration=configuration,
                responsible=responsible, serial=serial,
                department_id=department_id, user=user, status=status,
                price=price, purchase_date=purchase_date, location=location,
            )
            detail = oplog.diff({}, _asset_snapshot(created), ASSET_FIELD_LABELS)
            oplog.log(request, 'asset', 'create', target=_asset_label(created),
                      detail=f'新增资产；{detail}' if detail else '新增资产')
            messages.success(request, f'资产 {name} 新增成功。')
            return redirect('library')
    context = {
        'departments': Department.objects.all(),
        'department_options': [{'id': d.id, 'text': d.name} for d in Department.objects.all()],
        'category_options': [{'id': c[0], 'text': c[0]} for c in Asset.CATEGORY_CHOICES],
        'status_options': [{'id': s[0], 'text': s[0]} for s in Asset.STATUS_CHOICES],
        'page_title': '新增资产',
        'active': 'library',
    }
    return render(request, 'asset_form.html', context)


@perm_required('manage_assets')
def asset_update(request, pk):
    asset = get_object_or_404(Asset, pk=pk)
    if request.method == 'POST':
        # 变更前快照：用于生成"字段：旧值 → 新值"的操作日志明细
        before = _asset_snapshot(asset)
        asset.name = request.POST.get('name', asset.name).strip()
        asset.category = request.POST.get('category', asset.category)
        asset.specification = request.POST.get('specification', '').strip()
        asset.configuration = request.POST.get('configuration', '').strip()
        asset.responsible = request.POST.get('responsible', '').strip()
        asset.serial = request.POST.get('serial', '').strip()
        asset.department_id = request.POST.get('department') or None
        asset.user = request.POST.get('user', '').strip()
        asset.status = request.POST.get('status', asset.status)
        asset.location = request.POST.get('location', '').strip()
        # 价格与购置日期允许被清空（原实现会保留旧值，导致"清不掉"）
        price_value = _parse_price(request.POST.get('price'))
        if price_value is not None:
            asset.price = price_value
        raw_date = (request.POST.get('purchase_date') or '').strip()
        asset.purchase_date = _parse_date(raw_date) if raw_date else None

        missing = _missing_required_asset(dict(
            asset_id=asset.asset_id, name=asset.name, category=asset.category,
            department=asset.department_id, responsible=asset.responsible,
            user=asset.user, location=asset.location, status=asset.status,
        ))
        if missing:
            messages.error(request, '请填写必填项：' + '、'.join(missing))
        elif price_value is None:
            messages.error(request, '资产价值必须是数字。')
        else:
            asset.save()
            # 刷新以拿到最新的关联部门名称，再做前后对比（若字段无变化则不写日志）
            asset.refresh_from_db()
            diff_text = oplog.diff(before, _asset_snapshot(asset), ASSET_FIELD_LABELS)
            if diff_text:
                oplog.log(request, 'asset', 'update', target=_asset_label(asset), detail=diff_text)
            messages.success(request, f'资产 {asset.name} 更新成功。')
            return redirect('library')
    context = {'asset': asset, 'departments': Department.objects.all(),
               'department_options': [{'id': d.id, 'text': d.name} for d in Department.objects.all()],
               'category_options': [{'id': c[0], 'text': c[0]} for c in Asset.CATEGORY_CHOICES],
               'status_options': [{'id': s[0], 'text': s[0]} for s in Asset.STATUS_CHOICES],
               'attachments': asset.attachments.all(),
               'attachment_accept': ','.join(AssetAttachment.ALLOWED_EXT),
               'attachment_max_mb': AssetAttachment.MAX_SIZE // (1024 * 1024),
               'page_title': '编辑资产', 'active': 'library'}
    return render(request, 'asset_form.html', context)


@perm_required('manage_assets')
def asset_delete(request, pk):
    asset = get_object_or_404(Asset, pk=pk)
    if request.method == 'POST':
        # 级联保护：资产下若有关联历史(领用/变更)，禁止直接删除，避免静默丢失审计记录
        n_req = asset.requisitions.count()
        n_chg = asset.changes.count()
        if n_req or n_chg:
            parts = []
            if n_req:
                parts.append(f'{n_req} 条领用记录')
            if n_chg:
                parts.append(f'{n_chg} 条变更记录')
            messages.error(
                request,
                f'资产 {asset.name} 关联有 {("、".join(parts))}，无法直接删除。'
                '请先删除或转移其关联的领用/变更记录后再操作。',
            )
        else:
            label = _asset_label(asset)
            asset.delete()
            oplog.log(request, 'asset', 'delete', target=label, detail='删除资产及其全部信息')
            messages.success(request, f'资产 {asset.name} 已删除。')
    return redirect('library')


@perm_required('manage_assets')
def asset_import(request):
    result = None
    if request.method == 'POST':
        upload = request.FILES.get('file')
        if not upload:
            messages.error(request, '请选择要上传的 Excel 文件。')
            return redirect('asset_import')
        if not upload.name.lower().endswith(('.xlsx', '.xls')):
            messages.error(request, '仅支持 .xlsx / .xls 格式的文件。')
            return redirect('asset_import')
        try:
            wb = load_workbook(upload, data_only=True)
        except Exception as e:
            messages.error(request, f'文件解析失败：{e}')
            return redirect('asset_import')

        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            messages.error(request, 'Excel 内容为空。')
            return redirect('asset_import')

        # 表头映射
        header_map = {}
        for idx, h in enumerate(rows[0]):
            if h is None:
                continue
            key = str(h).strip()
            for field, names in _HEADER_ALIASES.items():
                if key in names:
                    header_map[field] = idx
                    break
        if 'asset_id' not in header_map or 'name' not in header_map:
            messages.error(request, '表格需至少包含「资产编号」与「资产名称」列。')
            return redirect('asset_import')

        def val(row, field):
            i = header_map.get(field)
            if i is None or i >= len(row):
                return ''
            v = row[i]
            return '' if v is None else v

        valid_status = [c[0] for c in Asset.STATUS_CHOICES]
        valid_cat = [c[0] for c in Asset.CATEGORY_CHOICES]
        department_cache = {d.name: d for d in Department.objects.all()}

        imported = skipped = 0
        warnings = []
        # 整批导入置于同一事务：中途异常即全部回滚，避免部分写入造成脏台账
        from django.db import transaction
        try:
            with transaction.atomic():
                for rn, row in enumerate(rows[1:], start=2):
                    if not row or all(str(c).strip() == '' for c in row if c is not None):
                        continue
                    asset_id = str(val(row, 'asset_id')).strip()
                    name = str(val(row, 'name')).strip()
                    if not asset_id or not name:
                        warnings.append(f'第 {rn} 行：资产编号或名称为空，已跳过')
                        continue
                    if Asset.objects.filter(asset_id=asset_id).exists():
                        skipped += 1
                        continue

                    # 类别：先按原值校验，不合法再查别名映射表归一化（如 笔记本→笔记本电脑、服务器→网络设备）
                    raw_cat = str(val(row, 'category')).strip()
                    category = raw_cat or '其他'
                    if category not in valid_cat:
                        mapped = Asset.CATEGORY_ALIAS.get(raw_cat)
                        if mapped and mapped in valid_cat:
                            warnings.append(f'{asset_id}：类别「{raw_cat}」已归一化为「{mapped}」')
                            category = mapped
                        else:
                            warnings.append(f'{asset_id}：类别「{raw_cat}」无法识别，已设为「其他」')
                            category = '其他'
                    # 状态：同样先原值校验，再别名归一化（如 闲置→库存、待报废→报废）
                    raw_status = str(val(row, 'status')).strip()
                    status = raw_status or '库存'
                    if status not in valid_status:
                        mapped_s = Asset.STATUS_ALIAS.get(raw_status)
                        if mapped_s and mapped_s in valid_status:
                            warnings.append(f'{asset_id}：状态「{raw_status}」已归一化为「{mapped_s}」')
                            status = mapped_s
                        else:
                            warnings.append(f'{asset_id}：状态「{raw_status}」无法识别，已设为「库存」')
                            status = '库存'

                    dept_name = str(val(row, 'department')).strip()
                    dept = department_cache.get(dept_name) if dept_name else None
                    if not dept:
                        warnings.append(f'{asset_id}：所属部门「{dept_name}」不存在或为空，已跳过')
                        continue

                    responsible = str(val(row, 'responsible')).strip()
                    user = str(val(row, 'user')).strip()
                    location = str(val(row, 'location')).strip()
                    if not responsible or not user or not location:
                        warnings.append(f'{asset_id}：责任人/实际使用人/存放位置不能为空，已跳过')
                        continue

                    price = 0
                    try:
                        pv = val(row, 'price')
                        price = float(pv) if pv not in ('', None) else 0
                    except Exception:
                        price = 0
                    purchase_date = _parse_date(val(row, 'purchase_date'))

                    Asset.objects.create(
                        asset_id=asset_id,
                        name=name,
                        category=category,
                        specification=str(val(row, 'specification')).strip(),
                        configuration=str(val(row, 'configuration')).strip(),
                        responsible=responsible,
                        serial=str(val(row, 'serial')).strip(),
                        user=user,
                        department=dept,
                        status=status,
                        price=price,
                        purchase_date=purchase_date,
                        location=location,
                    )
                    imported += 1
        except Exception as e:
            # 记录错误，但事务已回滚，不落任何半套数据
            result = {'imported': 0, 'skipped': 0, 'warnings': [], 'error': str(e)}
            messages.error(request, f'导入失败，已全部回滚：{e}')
            context = {'page_title': '导入资产', 'active': 'library', 'result': result}
            return render(request, 'asset_import.html', context)

        result = {'imported': imported, 'skipped': skipped, 'warnings': warnings}
        msg = f'导入完成：成功 {imported} 条'
        if skipped:
            msg += f'，跳过（编号已存在）{skipped} 条'
        if warnings:
            msg += f'，{len(warnings)} 处提示'
        oplog.log(
            request, 'asset', 'import', target=upload.name,
            detail=(f'成功导入 {imported} 条，编号已存在跳过 {skipped} 条'
                    + (f'，{len(warnings)} 处提示：' + '；'.join(warnings[:20]) if warnings else '')),
        )
        (messages.warning if warnings else messages.success)(request, msg)

    context = {'page_title': '导入资产', 'active': 'library', 'result': result}
    return render(request, 'asset_import.html', context)


@perm_required('view_assets')
def asset_import_template(request):
    wb = Workbook()
    ws = wb.active
    ws.title = '资产导入模板'
    headers = [
        '资产编号', '资产名称', '资产类别', '所属部门', '责任人', '实际使用人',
        '存放位置', '使用状态', '序列号', '规格', '配置', '价格', '购置日期',
    ]
    ws.append(headers)
    ws.append([
        'ZC-1001', '示例笔记本电脑', '电脑', '研发一部', '张伟', '张伟',
        'A-101', '在用', 'MP2HPWF', '14英寸', 'i7/16G/512G', 8999, '2023-01-01',
    ])
    widths = [12, 22, 12, 14, 12, 14, 16, 10, 16, 14, 20, 12, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="asset_import_template.xlsx"'
    wb.save(response)
    return response

