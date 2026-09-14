import csv
import datetime
import hashlib
import math
import os
from decimal import Decimal, InvalidOperation
from functools import wraps
from urllib.parse import urlencode

from openpyxl import Workbook
from openpyxl import load_workbook

from django.contrib import messages
from django.contrib.auth import authenticate, login, logout as auth_logout, update_session_auth_hash
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.core.cache import cache
from django.core.paginator import Paginator
from django.db.models import Count, Q, Sum
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.utils import timezone
from django.utils.cache import get_conditional_response
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_GET
from . import asset_snapshot as snapshot
from . import oplog
from .models import (
    Department, Asset, Requisition, AssetChange, Role, UserProfile,
    AssetAttachment, OperationLog,
)
from .permissions import perm_required
from .qr import qr_svg

# 登录安全：连续输错密码锁定
LOGIN_MAX_ATTEMPTS = 5        # 连续输错次数上限
LOGIN_LOCK_SECONDS = 5 * 60   # 锁定时长（秒）


def _login_fail_key(username):
    return 'login_fail_' + (username or '').strip().lower()


# 资产编辑留痕：参与"字段级变更明细"的字段与中文标签
ASSET_FIELD_LABELS = {
    'name': '资产名称',
    'category': '资产类别',
    'specification': '规格',
    'configuration': '配置',
    'responsible': '责任人',
    'serial': '序列号',
    'department': '所属部门',
    'user': '实际使用人',
    'status': '使用状态',
    'price': '资产价值',
    'purchase_date': '购置日期',
    'location': '存放位置',
}


def _asset_snapshot(asset):
    """把资产关键字段拍成 {字段名: 值} 快照，供变更前后对比。"""
    return {
        'name': asset.name,
        'category': asset.category,
        'specification': asset.specification,
        'configuration': asset.configuration,
        'responsible': asset.responsible,
        'serial': asset.serial,
        'department': asset.department.name if asset.department_id else '',
        'user': asset.user,
        'status': asset.status,
        'price': asset.price,
        'purchase_date': asset.purchase_date.strftime('%Y-%m-%d') if asset.purchase_date else '',
        'location': asset.location,
    }


def _asset_label(asset):
    """日志里标识一台资产的统一写法。"""
    return f'{asset.asset_id} {asset.name}'


# Excel 表头别名映射（兼容不同写法）
_HEADER_ALIASES = {
    'asset_id': ['资产编号', '编号', '资产ID'],
    'name': ['资产名称', '名称'],
    'category': ['类别', '资产类别'],
    'specification': ['规格', '型号'],
    'configuration': ['配置', '配置参数'],
    'department': ['所属部门', '部门', '责任部门'],
    'user': ['实际使用人', '使用人'],
    'responsible': ['责任人', '责任人员'],
    'serial': ['序列号', '机身序列号', '序列号(SN)'],
    'status': ['状态', '使用状态', '资产状态'],
    'price': ['价格', '价值', '资产价值', '价格(元)'],
    'purchase_date': ['购置日期', '日期', '购置时间'],
    'location': ['位置', '存放位置', '存放地点'],
}


def _parse_date(value):
    """兼容字符串日期 / datetime / 数值日期，返回 date 或 None。"""
    if value is None or str(value).strip() == '':
        return None
    if isinstance(value, datetime.datetime):
        return value.date() if hasattr(value, 'date') else None
    if isinstance(value, datetime.date):
        return value
    from openpyxl.utils.datetime import from_excel
    try:
        return from_excel(float(value)).date()
    except Exception:
        pass
    text = str(value).strip().replace('/', '-').replace('.', '-')
    for fmt in ('%Y-%m-%d', '%Y-%m-%d %H:%M:%S', '%Y年%m月%d日'):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except Exception:
            continue
    return None


def _parse_price(value):
    """把价格输入解析为 Decimal：空值视为 0，非法输入返回 None。"""
    text = str(value).strip() if value is not None else ''
    if not text:
        return Decimal('0')
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _check_password_strength(user, password):
    """按 settings.AUTH_PASSWORD_VALIDATORS 校验密码强度。

    通过返回 None，否则返回可直接展示的错误文案。
    """
    from django.contrib.auth.password_validation import validate_password
    from django.core.exceptions import ValidationError
    try:
        validate_password(password, user=user)
        return None
    except ValidationError as exc:
        return '；'.join(exc.messages)


def admin_required(view_func):
    """仅管理员(staff/superuser)可访问；否则提示并回首页。"""
    @wraps(view_func)
    @login_required
    def _wrapped(request, *args, **kwargs):
        if not (request.user.is_staff or request.user.is_superuser):
            messages.error(request, '仅管理员可访问该页面。')
            return redirect(reverse('dashboard'))
        return view_func(request, *args, **kwargs)
    return _wrapped


# ---------- 登录 ----------
def login_view(request):
    if request.user.is_authenticated:
        return redirect(reverse('dashboard'))

    kicked = request.GET.get('kicked') == '1'

    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '')
        remember = request.POST.get('remember')

        # —— 防爆破：连错 5 次锁定 5 分钟（按用户名计数，锁定期间即使密码正确也拒绝）——
        fail_key = _login_fail_key(username)
        lock_key = fail_key + '_lock'
        locked_until = cache.get(lock_key)
        if locked_until:
            remain = max(1, math.ceil((locked_until - timezone.now()).total_seconds() / 60))
            oplog.log(request, 'auth', 'login_fail', target=username or '（空用户名）',
                      detail=f'账号处于锁定状态，拒绝登录（剩余约 {remain} 分钟）')
            return render(request, 'login.html', {
                'error': True,
                'error_msg': f'密码连续输错 {LOGIN_MAX_ATTEMPTS} 次，账号已临时锁定，请约 {remain} 分钟后再试。',
                'username': username,
            })

        user = authenticate(request, username=username, password=password)
        if user is not None:
            cache.delete(fail_key)
            cache.delete(lock_key)
            login(request, user)
            # 记住我：30 天；否则浏览器关闭即失效
            if remember:
                request.session.set_expiry(60 * 60 * 24 * 30)
            else:
                request.session.set_expiry(0)
            # —— 单设备登录：非超管只保留最近一次登录的会话，旧设备由中间件踢下线 ——
            profile = getattr(user, 'profile', None)
            if profile is not None and not user.is_superuser:
                profile.session_key = request.session.session_key
                profile.save(update_fields=['session_key'])
            oplog.log(request, 'auth', 'login', target=user.username,
                      detail=f'登录成功（{"记住我 30 天" if remember else "关闭浏览器即失效"}）')
            nxt = request.POST.get('next') or request.GET.get('next') or '/'
            if not url_has_allowed_host_and_scheme(
                nxt, allowed_hosts={request.get_host()}, require_https=request.is_secure()
            ):
                nxt = '/'
            return redirect(nxt)

        # 登录失败：计数并提示剩余机会
        fails = cache.get(fail_key, 0) + 1
        if fails >= LOGIN_MAX_ATTEMPTS:
            cache.set(lock_key, timezone.now() + datetime.timedelta(seconds=LOGIN_LOCK_SECONDS),
                      LOGIN_LOCK_SECONDS)
            cache.delete(fail_key)
            error_msg = f'密码连续输错 {LOGIN_MAX_ATTEMPTS} 次，账号已锁定 5 分钟，请稍后再试。'
            oplog.log(request, 'auth', 'login_fail', target=username or '（空用户名）',
                      detail=f'连续输错 {LOGIN_MAX_ATTEMPTS} 次，账号锁定 {LOGIN_LOCK_SECONDS // 60} 分钟')
        else:
            cache.set(fail_key, fails, LOGIN_LOCK_SECONDS)
            error_msg = (f'用户名或密码不正确，请重试。'
                         f'（第 {fails}/{LOGIN_MAX_ATTEMPTS} 次尝试，连续输错 {LOGIN_MAX_ATTEMPTS} 次将锁定 5 分钟）')
            oplog.log(request, 'auth', 'login_fail', target=username or '（空用户名）',
                      detail=f'用户名或密码不正确（第 {fails}/{LOGIN_MAX_ATTEMPTS} 次）')
        return render(request, 'login.html', {
            'error': True, 'error_msg': error_msg, 'username': username,
        })

    return render(request, 'login.html', {'kicked': kicked})


def logout_view(request):
    """退出登录（记录日志后走 Django 的 logout）。"""
    if request.user.is_authenticated:
        oplog.log(request, 'auth', 'logout', target=request.user.username, detail='主动退出登录')
    auth_logout(request)
    return redirect('login')


@login_required
def password_change(request):
    if request.method == 'POST':
        old_password = request.POST.get('old_password', '')
        new_password = request.POST.get('new_password', '')
        confirm_password = request.POST.get('confirm_password', '')
        if not request.user.check_password(old_password):
            messages.error(request, '当前密码不正确。')
        elif new_password != confirm_password:
            messages.error(request, '两次输入的新密码不一致。')
        else:
            pwd_err = _check_password_strength(request.user, new_password)
            if pwd_err:
                messages.error(request, f'新密码不符合要求：{pwd_err}')
            else:
                request.user.set_password(new_password)
                request.user.save()
                update_session_auth_hash(request, request.user)
                oplog.log(request, 'auth', 'update', target=request.user.username,
                          detail='用户自行修改登录密码')
                messages.success(request, '密码修改成功。')
                return redirect('dashboard')
    context = {'page_title': '修改密码', 'active': 'dashboard'}
    return render(request, 'password_change.html', context)


# ---------- 首页仪表盘 ----------
@login_required
def dashboard(request):
    # 各部门资产数量统计（用于饼图占比 与 柱状图）
    dept_counts = (
        Department.objects.annotate(asset_count=Count('assets'))
        .order_by('-asset_count')
    )
    labels = [d.name for d in dept_counts]
    counts = [d.asset_count for d in dept_counts]

    # 状态分布（饼图补充）
    status_counts = (
        Asset.objects.values('status')
        .annotate(n=Count('id'))
        .order_by('-n')
    )

    # 类别分布
    category_counts = (
        Asset.objects.values('category')
        .annotate(n=Count('id'))
        .order_by('-n')
    )

    # 各部门资产总值
    value_per_dept = (
        Department.objects.annotate(total_value=Sum('assets__price'))
        .filter(total_value__isnull=False)
        .exclude(total_value=0)
        .order_by('-total_value')
    )
    value_labels = [d.name for d in value_per_dept]
    value_counts = [float(d.total_value or 0) for d in value_per_dept]

    total_assets = Asset.objects.count()
    in_use = Asset.objects.filter(status='在用').count()
    total_departments = Department.objects.count()
    # 逾期为派生状态（借出且超过预计归还日期），用与列表页一致的动态口径统计
    pending_return = Requisition.objects.filter(status='借出').count()
    overdue = Requisition.objects.filter(
        status='借出', due_date__lt=timezone.localdate()
    ).count()
    # 逾期清单（按超期天数倒序），供首页「逾期提醒」卡片展示与跳转
    overdue_records = list(
        Requisition.objects.filter(status='借出', due_date__lt=timezone.localdate())
        .select_related('asset', 'department')
        .order_by('due_date')[:8]
    )
    for r in overdue_records:
        r.overdue_days = (timezone.localdate() - r.due_date).days
    total_value = Asset.objects.aggregate(v=Sum('price'))['v'] or 0

    context = {
        'labels': labels,
        'counts': counts,
        'status_labels': [s['status'] for s in status_counts],
        'status_values': [s['n'] for s in status_counts],
        'category_labels': [c['category'] for c in category_counts],
        'category_values': [c['n'] for c in category_counts],
        'value_labels': value_labels,
        'value_counts': value_counts,
        'total_assets': total_assets,
        'in_use': in_use,
        'total_departments': total_departments,
        'overdue': overdue,
        'overdue_records': overdue_records,
        'total_value': total_value,
        'pending_return': pending_return,
        'recent_changes': AssetChange.objects.select_related('asset')[:6],
        'recent_requisitions': Requisition.objects.select_related('asset', 'department')[:6],
        'page_title': '数据总览',
    }
    return render(request, 'dashboard.html', context)


def _filtered_assets(request):
    """按 URL 查询参数过滤资产库，返回 (queryset, filters)。

    资产库列表与「标签打印」共用同一套筛选口径，避免两处规则漂移。
    """
    status = request.GET.get('status', '')
    category = request.GET.get('category', '')
    dept_id = request.GET.get('department', '')
    keyword = request.GET.get('q', '').strip()
    # 搜索模式：asset=按名称/资产编号；person=按责任人/实际使用人
    field = request.GET.get('field', 'asset')
    if field not in ('asset', 'person'):
        field = 'asset'

    qs = Asset.objects.select_related('department')
    if status:
        qs = qs.filter(status=status)
    if category:
        qs = qs.filter(category=category)
    if dept_id:
        qs = qs.filter(department_id=dept_id)
    if keyword:
        if field == 'person':
            # 责任人 / 实际使用人 模糊匹配
            qs = qs.filter(Q(responsible__icontains=keyword) | Q(user__icontains=keyword))
        else:
            # 资产名称 / 资产编号 模糊匹配（默认）
            qs = qs.filter(Q(name__icontains=keyword) | Q(asset_id__icontains=keyword))

    filters = {
        'status': status,
        'category': category,
        'department': dept_id,
        'q': keyword,
        'field': field,
    }
    return qs, filters


# ---------- 资产库模块 ----------
@perm_required('view_assets')
def library(request):
    qs, filters = _filtered_assets(request)
    status = filters['status']
    category = filters['category']
    dept_id = filters['department']
    keyword = filters['q']
    field = filters['field']
    sort = request.GET.get('sort', 'asset_id')
    direction = request.GET.get('dir', 'asc')

    # 排序白名单
    sort_map = {
        'asset_id': 'asset_id', 'name': 'name', 'category': 'category',
        'price': 'price', 'status': 'status', 'department': 'department__name',
    }
    order_field = sort_map.get(sort, 'asset_id')
    order = ('-' if direction == 'desc' else '') + order_field
    qs = qs.order_by(order)

    # 导出 CSV
    if request.GET.get('export') == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response.write('\ufeff')  # BOM，方便 Excel 识别中文
        w = csv.writer(response)
        w.writerow(['资产编号', '资产名称', '资产类别', '所属部门', '责任人', '实际使用人', '存放位置', '使用状态', '序列号', '规格', '配置', '价格(元)'])
        for a in qs:
            w.writerow([
                a.asset_id, a.name, a.category, a.department.name if a.department else '',
                a.responsible, a.user, a.location, a.status, a.serial,
                a.specification, a.configuration, a.price,
            ])
        filename = 'assets_export.csv'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    # 每页条数
    try:
        per_page = int(request.GET.get('per_page', 12))
    except (TypeError, ValueError):
        per_page = 12
    if per_page not in (10, 20, 50, 100):
        per_page = 12

    paginator = Paginator(qs, per_page)
    page_obj = paginator.get_page(request.GET.get('page'))

    # 保留筛选参数的查询串（排除排序/方向/页码），供排序链接复用
    keep = {k: v for k, v in request.GET.items() if k not in ('sort', 'dir', 'page')}
    base_qs = urlencode(keep)
    qs_no_perpage = urlencode({k: v for k, v in request.GET.items() if k not in ('page', 'per_page')})

    context = {
        'page_obj': page_obj,
        'assets': page_obj.object_list,
        'departments': Department.objects.all(),
        'category_options': [c[0] for c in Asset.CATEGORY_CHOICES],
        'status_options': [s[0] for s in Asset.STATUS_CHOICES],
        'page_title': '资产库',
        'active': 'library',
        'sort': sort,
        'direction': direction,
        'base_qs': base_qs,
        'qs_no_perpage': qs_no_perpage,
        'per_page': per_page,
        'per_page_choices': (10, 20, 50, 100),
        'filters': filters,
    }
    return render(request, 'library.html', context)


@perm_required('manage_assets')
def assets_bulk_delete(request):
    if request.method == 'POST':
        idstr = request.POST.get('ids', '').strip()
        ids = [int(x) for x in idstr.split(',') if x.strip().isdigit()]
        if ids:
            # 级联保护：跳过关联有领用/变更历史的资产，避免静默丢失审计记录
            protected = []
            deletable = []
            for a in Asset.objects.filter(pk__in=ids).select_related('department'):
                if a.requisitions.exists() or a.changes.exists():
                    protected.append(a.asset_id)
                else:
                    deletable.append(a.pk)
            deleted_n = len(deletable)
            if deleted_n:
                Asset.objects.filter(pk__in=deletable).delete()
            if deleted_n or protected:
                detail_parts = [f'已删除 {deleted_n} 项']
                if protected:
                    detail_parts.append(
                        f'因存在领用/变更历史被跳过 {len(protected)} 项：'
                        f'{", ".join(protected[:20])}'
                    )
                oplog.log(request, 'asset', 'delete',
                          target=f'批量删除（{deleted_n} 项）',
                          detail='；'.join(detail_parts))
            if protected:
                messages.warning(
                    request,
                    f'已跳过 {len(protected)} 项存在关联领用/变更历史的资产'
                    f'（{", ".join(protected[:8])}{"…" if len(protected) > 8 else ""}），请逐项处理后删除。',
                )
            if deleted_n:
                messages.success(request, f'已删除 {deleted_n} 条资产。')
            if not deleted_n and not protected:
                messages.error(request, '没有可删除的资产。')
        else:
            messages.error(request, '请先勾选要删除的资产。')
        nxt = request.POST.get('next') or reverse('library')
        return redirect(nxt)
    messages.error(request, '非法请求。')
    return redirect(reverse('library'))


@perm_required('manage_assets')
def asset_create(request):
    if request.method == 'POST':
        asset_id = request.POST.get('asset_id', '').strip()
        name = request.POST.get('name', '').strip()
        category = request.POST.get('category', '')
        specification = request.POST.get('specification', '').strip()
        configuration = request.POST.get('configuration', '').strip()
        responsible = request.POST.get('responsible', '').strip()
        serial = request.POST.get('serial', '').strip()
        department_id = request.POST.get('department') or None
        user = request.POST.get('user', '').strip()
        status = request.POST.get('status', '库存')
        price_value = _parse_price(request.POST.get('price'))
        price = price_value if price_value is not None else 0
        raw_date = (request.POST.get('purchase_date') or '').strip()
        purchase_date = _parse_date(raw_date) if raw_date else None
        location = request.POST.get('location', '').strip()

        missing = _missing_required_asset(dict(
            asset_id=asset_id, name=name, category=category, department=department_id,
            responsible=responsible, user=user, location=location, status=status,
        ))
        if missing:
            messages.error(request, '请填写必填项：' + '、'.join(missing))
        elif Asset.objects.filter(asset_id=asset_id).exists():
            messages.error(request, f'资产编号 {asset_id} 已存在。')
        elif price_value is None:
            messages.error(request, '资产价值必须是数字。')
        else:
            created = Asset.objects.create(
                asset_id=asset_id, name=name, category=category,
                specification=specification, configuration=configuration,
                responsible=responsible, serial=serial,
                department_id=department_id, user=user, status=status,
                price=price, purchase_date=purchase_date, location=location,
            )
            detail = oplog.diff({}, _asset_snapshot(created), ASSET_FIELD_LABELS)
            oplog.log(request, 'asset', 'create', target=_asset_label(created),
                      detail=f'新增资产；{detail}' if detail else '新增资产')
            messages.success(request, f'资产 {name} 新增成功。')
            return redirect('library')
    context = {
        'departments': Department.objects.all(),
        'department_options': [{'id': d.id, 'text': d.name} for d in Department.objects.all()],
        'category_options': [{'id': c[0], 'text': c[0]} for c in Asset.CATEGORY_CHOICES],
        'status_options': [{'id': s[0], 'text': s[0]} for s in Asset.STATUS_CHOICES],
        'page_title': '新增资产',
        'active': 'library',
    }
    return render(request, 'asset_form.html', context)


def _missing_required_asset(data):
    labels = {
        'asset_id': '资产编号', 'name': '资产名称', 'category': '资产类别',
        'department': '所属部门', 'responsible': '责任人', 'user': '实际使用人',
        'location': '存放位置', 'status': '使用状态',
    }
    return [labels[k] for k, v in data.items() if v in (None, '')]


@perm_required('manage_assets')
def asset_update(request, pk):
    asset = get_object_or_404(Asset, pk=pk)
    if request.method == 'POST':
        # 变更前快照：用于生成"字段：旧值 → 新值"的操作日志明细
        before = _asset_snapshot(asset)
        asset.name = request.POST.get('name', asset.name).strip()
        asset.category = request.POST.get('category', asset.category)
        asset.specification = request.POST.get('specification', '').strip()
        asset.configuration = request.POST.get('configuration', '').strip()
        asset.responsible = request.POST.get('responsible', '').strip()
        asset.serial = request.POST.get('serial', '').strip()
        asset.department_id = request.POST.get('department') or None
        asset.user = request.POST.get('user', '').strip()
        asset.status = request.POST.get('status', asset.status)
        asset.location = request.POST.get('location', '').strip()
        # 价格与购置日期允许被清空（原实现会保留旧值，导致"清不掉"）
        price_value = _parse_price(request.POST.get('price'))
        if price_value is not None:
            asset.price = price_value
        raw_date = (request.POST.get('purchase_date') or '').strip()
        asset.purchase_date = _parse_date(raw_date) if raw_date else None

        missing = _missing_required_asset(dict(
            asset_id=asset.asset_id, name=asset.name, category=asset.category,
            department=asset.department_id, responsible=asset.responsible,
            user=asset.user, location=asset.location, status=asset.status,
        ))
        if missing:
            messages.error(request, '请填写必填项：' + '、'.join(missing))
        elif price_value is None:
            messages.error(request, '资产价值必须是数字。')
        else:
            asset.save()
            # 刷新以拿到最新的关联部门名称，再做前后对比（若字段无变化则不写日志）
            asset.refresh_from_db()
            diff_text = oplog.diff(before, _asset_snapshot(asset), ASSET_FIELD_LABELS)
            if diff_text:
                oplog.log(request, 'asset', 'update', target=_asset_label(asset), detail=diff_text)
            messages.success(request, f'资产 {asset.name} 更新成功。')
            return redirect('library')
    context = {'asset': asset, 'departments': Department.objects.all(),
               'department_options': [{'id': d.id, 'text': d.name} for d in Department.objects.all()],
               'category_options': [{'id': c[0], 'text': c[0]} for c in Asset.CATEGORY_CHOICES],
               'status_options': [{'id': s[0], 'text': s[0]} for s in Asset.STATUS_CHOICES],
               'attachments': asset.attachments.all(),
               'attachment_accept': ','.join(AssetAttachment.ALLOWED_EXT),
               'attachment_max_mb': AssetAttachment.MAX_SIZE // (1024 * 1024),
               'page_title': '编辑资产', 'active': 'library'}
    return render(request, 'asset_form.html', context)


@perm_required('manage_assets')
def asset_delete(request, pk):
    asset = get_object_or_404(Asset, pk=pk)
    if request.method == 'POST':
        # 级联保护：资产下若有关联历史(领用/变更)，禁止直接删除，避免静默丢失审计记录
        n_req = asset.requisitions.count()
        n_chg = asset.changes.count()
        if n_req or n_chg:
            parts = []
            if n_req:
                parts.append(f'{n_req} 条领用记录')
            if n_chg:
                parts.append(f'{n_chg} 条变更记录')
            messages.error(
                request,
                f'资产 {asset.name} 关联有 {("、".join(parts))}，无法直接删除。'
                '请先删除或转移其关联的领用/变更记录后再操作。',
            )
        else:
            label = _asset_label(asset)
            asset.delete()
            oplog.log(request, 'asset', 'delete', target=label, detail='删除资产及其全部信息')
            messages.success(request, f'资产 {asset.name} 已删除。')
    return redirect('library')


# ---------- 资产 Excel 导入 ----------
@perm_required('manage_assets')
def asset_import(request):
    result = None
    if request.method == 'POST':
        upload = request.FILES.get('file')
        if not upload:
            messages.error(request, '请选择要上传的 Excel 文件。')
            return redirect('asset_import')
        if not upload.name.lower().endswith(('.xlsx', '.xls')):
            messages.error(request, '仅支持 .xlsx / .xls 格式的文件。')
            return redirect('asset_import')
        try:
            wb = load_workbook(upload, data_only=True)
        except Exception as e:
            messages.error(request, f'文件解析失败：{e}')
            return redirect('asset_import')

        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            messages.error(request, 'Excel 内容为空。')
            return redirect('asset_import')

        # 表头映射
        header_map = {}
        for idx, h in enumerate(rows[0]):
            if h is None:
                continue
            key = str(h).strip()
            for field, names in _HEADER_ALIASES.items():
                if key in names:
                    header_map[field] = idx
                    break
        if 'asset_id' not in header_map or 'name' not in header_map:
            messages.error(request, '表格需至少包含「资产编号」与「资产名称」列。')
            return redirect('asset_import')

        def val(row, field):
            i = header_map.get(field)
            if i is None or i >= len(row):
                return ''
            v = row[i]
            return '' if v is None else v

        valid_status = [c[0] for c in Asset.STATUS_CHOICES]
        valid_cat = [c[0] for c in Asset.CATEGORY_CHOICES]
        department_cache = {d.name: d for d in Department.objects.all()}

        imported = skipped = 0
        warnings = []
        # 整批导入置于同一事务：中途异常即全部回滚，避免部分写入造成脏台账
        from django.db import transaction
        try:
            with transaction.atomic():
                for rn, row in enumerate(rows[1:], start=2):
                    if not row or all(str(c).strip() == '' for c in row if c is not None):
                        continue
                    asset_id = str(val(row, 'asset_id')).strip()
                    name = str(val(row, 'name')).strip()
                    if not asset_id or not name:
                        warnings.append(f'第 {rn} 行：资产编号或名称为空，已跳过')
                        continue
                    if Asset.objects.filter(asset_id=asset_id).exists():
                        skipped += 1
                        continue

                    # 类别：先按原值校验，不合法再查别名映射表归一化（如 笔记本→笔记本电脑、服务器→网络设备）
                    raw_cat = str(val(row, 'category')).strip()
                    category = raw_cat or '其他'
                    if category not in valid_cat:
                        mapped = Asset.CATEGORY_ALIAS.get(raw_cat)
                        if mapped and mapped in valid_cat:
                            warnings.append(f'{asset_id}：类别「{raw_cat}」已归一化为「{mapped}」')
                            category = mapped
                        else:
                            warnings.append(f'{asset_id}：类别「{raw_cat}」无法识别，已设为「其他」')
                            category = '其他'
                    # 状态：同样先原值校验，再别名归一化（如 闲置→库存、待报废→报废）
                    raw_status = str(val(row, 'status')).strip()
                    status = raw_status or '库存'
                    if status not in valid_status:
                        mapped_s = Asset.STATUS_ALIAS.get(raw_status)
                        if mapped_s and mapped_s in valid_status:
                            warnings.append(f'{asset_id}：状态「{raw_status}」已归一化为「{mapped_s}」')
                            status = mapped_s
                        else:
                            warnings.append(f'{asset_id}：状态「{raw_status}」无法识别，已设为「库存」')
                            status = '库存'

                    dept_name = str(val(row, 'department')).strip()
                    dept = department_cache.get(dept_name) if dept_name else None
                    if not dept:
                        warnings.append(f'{asset_id}：所属部门「{dept_name}」不存在或为空，已跳过')
                        continue

                    responsible = str(val(row, 'responsible')).strip()
                    user = str(val(row, 'user')).strip()
                    location = str(val(row, 'location')).strip()
                    if not responsible or not user or not location:
                        warnings.append(f'{asset_id}：责任人/实际使用人/存放位置不能为空，已跳过')
                        continue

                    price = 0
                    try:
                        pv = val(row, 'price')
                        price = float(pv) if pv not in ('', None) else 0
                    except Exception:
                        price = 0
                    purchase_date = _parse_date(val(row, 'purchase_date'))

                    Asset.objects.create(
                        asset_id=asset_id,
                        name=name,
                        category=category,
                        specification=str(val(row, 'specification')).strip(),
                        configuration=str(val(row, 'configuration')).strip(),
                        responsible=responsible,
                        serial=str(val(row, 'serial')).strip(),
                        user=user,
                        department=dept,
                        status=status,
                        price=price,
                        purchase_date=purchase_date,
                        location=location,
                    )
                    imported += 1
        except Exception as e:
            # 记录错误，但事务已回滚，不落任何半套数据
            result = {'imported': 0, 'skipped': 0, 'warnings': [], 'error': str(e)}
            messages.error(request, f'导入失败，已全部回滚：{e}')
            context = {'page_title': '导入资产', 'active': 'library', 'result': result}
            return render(request, 'asset_import.html', context)

        result = {'imported': imported, 'skipped': skipped, 'warnings': warnings}
        msg = f'导入完成：成功 {imported} 条'
        if skipped:
            msg += f'，跳过（编号已存在）{skipped} 条'
        if warnings:
            msg += f'，{len(warnings)} 处提示'
        oplog.log(
            request, 'asset', 'import', target=upload.name,
            detail=(f'成功导入 {imported} 条，编号已存在跳过 {skipped} 条'
                    + (f'，{len(warnings)} 处提示：' + '；'.join(warnings[:20]) if warnings else '')),
        )
        (messages.warning if warnings else messages.success)(request, msg)

    context = {'page_title': '导入资产', 'active': 'library', 'result': result}
    return render(request, 'asset_import.html', context)


@perm_required('view_assets')
def asset_import_template(request):
    wb = Workbook()
    ws = wb.active
    ws.title = '资产导入模板'
    headers = [
        '资产编号', '资产名称', '资产类别', '所属部门', '责任人', '实际使用人',
        '存放位置', '使用状态', '序列号', '规格', '配置', '价格', '购置日期',
    ]
    ws.append(headers)
    ws.append([
        'ZC-1001', '示例笔记本电脑', '电脑', '研发一部', '张伟', '张伟',
        'A-101', '在用', 'MP2HPWF', '14英寸', 'i7/16G/512G', 8999, '2023-01-01',
    ])
    widths = [12, 22, 12, 14, 12, 14, 16, 10, 16, 14, 20, 12, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="asset_import_template.xlsx"'
    wb.save(response)
    return response


# ---------- 组织架构模块 ----------
@perm_required('view_org')
def org(request):
    mode = request.GET.get('mode', 'list')
    departments = Department.objects.annotate(asset_count=Count('assets')).order_by('id')

    # 汇总：每个部门统计其自身 + 所有下级部门的资产数（公司主体应包含下属全部）
    direct = {d.id: d.asset_count for d in departments}
    children = {}
    for d in departments:
        if d.parent_id:
            children.setdefault(d.parent_id, []).append(d.id)
    sub_memo = {}
    def subtree(nid):
        if nid not in sub_memo:
            total = direct.get(nid, 0)
            for cid in children.get(nid, []):
                total += subtree(cid)
            sub_memo[nid] = total
        return sub_memo[nid]
    for d in departments:
        d.subtree_count = subtree(d.id)

    # 构建树（仅树状模式需要）
    nodes = {d.id: {'dept': d, 'children': []} for d in departments}
    roots = []
    for d in departments:
        node = nodes[d.id]
        if d.parent_id and d.parent_id in nodes:
            nodes[d.parent_id]['children'].append(node)
        else:
            roots.append(node)

    context = {
        'departments': departments,
        'mode': mode,
        'tree_roots': roots,
        'total_departments': Department.objects.count(),
        'total_assets': Asset.objects.count(),
        'page_title': '组织架构',
        'active': 'org',
    }
    return render(request, 'org.html', context)


def _descendant_ids(dept):
    """返回某个部门所有下级部门的 id 列表（用于防止把上级设为下级造成成环）。"""
    ids = []
    stack = list(dept.children.all())
    while stack:
        child = stack.pop()
        ids.append(child.id)
        stack.extend(child.children.all())
    return ids


@perm_required('manage_org')
def department_form(request, pk=None):
    """新增/编辑部门（树状图与列表共用）。pk 为空则新增，否则编辑。"""
    is_edit = pk is not None
    department = get_object_or_404(Department, pk=pk) if is_edit else None
    preset_parent = request.GET.get('parent') or None
    mode = request.GET.get('mode', request.POST.get('mode', 'list'))

    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        manager = request.POST.get('manager', '').strip()
        description = request.POST.get('description', '').strip()
        parent_id = request.POST.get('parent') or None
        mode = request.POST.get('mode', mode)

        # 校验是否成环：编辑时不能把上级设为自身或其下级
        cycle = False
        if is_edit and parent_id:
            pid = int(parent_id)
            if pid == department.pk or pid in _descendant_ids(department):
                cycle = True

        dup = Department.objects.filter(name=name)
        if is_edit:
            dup = dup.exclude(pk=department.pk)

        if not name:
            messages.error(request, '部门名称不能为空。')
        elif dup.exists():
            messages.error(request, f'部门 {name} 已存在。')
        elif cycle:
            messages.error(request, '上级部门不能设为自身或其下级部门。')
        else:
            if is_edit:
                department.name = name
                department.manager = manager
                department.description = description
                department.parent_id = parent_id
                department.save()
                oplog.log(request, 'org', 'update', target=f'部门 {name}',
                          detail=(f'负责人：{manager or "（空）"}；'
                                  f'上级部门：{department.parent.name if department.parent_id else "（无）"}'))
                messages.success(request, f'部门 {name} 更新成功。')
            else:
                created_dept = Department.objects.create(
                    name=name, manager=manager, description=description,
                    parent_id=parent_id,
                )
                oplog.log(request, 'org', 'create', target=f'部门 {name}',
                          detail=(f'负责人：{manager or "（空）"}；'
                                  f'上级部门：{created_dept.parent.name if created_dept.parent_id else "（无）"}'))
                messages.success(request, f'部门 {name} 创建成功。')
            return redirect(f"{reverse('org')}?mode={mode}")

    context = {
        'department': department,
        'departments': Department.objects.all(),
        'is_edit': is_edit,
        'mode': mode,
        'preset_parent': int(preset_parent) if preset_parent else None,
        'page_title': '编辑部门' if is_edit else '新增部门',
        'active': 'org',
    }
    return render(request, 'department_form.html', context)


@perm_required('manage_org')
def department_delete(request, pk):
    department = get_object_or_404(Department, pk=pk)
    mode = request.GET.get('mode', 'list')
    if request.method == 'POST':
        mode = request.POST.get('mode', mode)
        if department.children.exists():
            messages.error(request, f'部门 {department.name} 下还有下级部门，请先移除或删除下级部门后再删除。')
        elif department.assets.exists():
            n = department.assets.count()
            messages.error(
                request,
                f'部门 {department.name} 下仍挂靠 {n} 项资产，删除将导致这些资产失去所属部门。'
                '请先将资产转移/分配到其它部门后再删除。',
            )
        else:
            from django.db import transaction
            with transaction.atomic():
                dept_name = department.name
                department.delete()
            oplog.log(request, 'org', 'delete', target=f'部门 {dept_name}', detail='删除部门')
            messages.success(request, f'部门 {department.name} 已删除。')
    return redirect(f"{reverse('org')}?mode={mode}")


# ---------- 资产领用模块 ----------
@perm_required('view_requisition')
def requisition(request):
    status = request.GET.get('status', '')

    all_records = list(Requisition.objects.select_related('asset', 'department').order_by('-borrow_date'))
    if status == '逾期':
        records = [r for r in all_records if r.display_status == '逾期']
    elif status:
        records = [r for r in all_records if r.status == status]
    else:
        records = all_records

    overdue_count = sum(1 for r in all_records if r.display_status == '逾期')

    paginator = Paginator(records, 10)
    page_obj = paginator.get_page(request.GET.get('page'))

    available_assets = Asset.objects.filter(status='库存')
    asset_options = [
        {'id': a.pk, 'asset_id': a.asset_id, 'name': a.name, 'status': a.status}
        for a in available_assets
    ]

    context = {
        'records': page_obj.object_list,
        'page_obj': page_obj,
        'assets': available_assets,
        'asset_options': asset_options,
        'department_options': [{'id': d.id, 'text': d.name} for d in Department.objects.all()],
        'departments': Department.objects.all(),
        'total_borrowed': Requisition.objects.filter(status='借出').count(),
        'total_returned': Requisition.objects.filter(status='已归还').count(),
        'overdue_count': overdue_count,
        'status': status,
        'page_title': '资产领用',
        'active': 'requisition',
    }
    return render(request, 'requisition.html', context)


@perm_required('view_requisition')
def requisition_export(request):
    """按当前筛选条件导出「资产领用」记录为 Excel。"""
    status = request.GET.get('status', '')
    all_records = list(Requisition.objects.select_related('asset', 'department').order_by('-borrow_date'))
    if status == '逾期':
        records = [r for r in all_records if r.display_status == '逾期']
    elif status:
        records = [r for r in all_records if r.status == status]
    else:
        records = all_records

    wb = Workbook()
    ws = wb.active
    ws.title = '资产领用'
    headers = ['资产编号', '资产名称', '领用人', '实际使用人', '领用部门', '领用用途',
               '领用日期', '预计归还', '归还日期', '状态']
    ws.append(headers)
    for r in records:
        ws.append([
            r.asset.asset_id,
            r.asset.name,
            r.user,
            r.actual_user or '',
            r.department.name if r.department else '',
            r.purpose,
            r.borrow_date.strftime('%Y-%m-%d') if r.borrow_date else '',
            r.due_date.strftime('%Y-%m-%d') if r.due_date else '',
            r.return_date.strftime('%Y-%m-%d') if r.return_date else '',
            r.display_status,
        ])
    for i, w in enumerate([14, 22, 12, 14, 14, 24, 12, 12, 12, 10], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="requisition_export.xlsx"'
    wb.save(response)
    return response


@perm_required('manage_requisition')
def requisition_create(request):
    if request.method == 'POST':
        asset_id = request.POST.get('asset')
        user = request.POST.get('user', '').strip()
        actual_user = request.POST.get('actual_user', '').strip()
        department_id = request.POST.get('department') or None
        purpose = request.POST.get('purpose', '').strip()
        due_date = request.POST.get('due_date') or None

        if not asset_id:
            messages.error(request, '请选择要领用的资产。')
        elif not user:
            messages.error(request, '领用人不能为空。')
        elif not department_id:
            messages.error(request, '请选择领用部门。')
        elif not due_date:
            messages.error(request, '请选择预计归还日期。')
        elif not purpose:
            messages.error(request, '请填写领用用途。')
        else:
            asset = get_object_or_404(Asset, pk=asset_id)
            if asset.status not in ['库存', '维修']:
                messages.error(request, f'资产 {asset.name} 当前状态为「{asset.status}」，不可领用。')
            else:
                Requisition.objects.create(
                    asset=asset, user=user, actual_user=actual_user,
                    department_id=department_id,
                    purpose=purpose, due_date=due_date,
                )
                oplog.log(
                    request, 'requisition', 'create', target=_asset_label(asset),
                    detail=(f'领用人：{user}；实际使用人：{actual_user or "同领用人"}；'
                            f'预计归还：{due_date}；用途：{purpose}'),
                )
                messages.success(
                    request,
                    f'资产 {asset.name} 已领用给 {user}，责任人、实际使用人与所属部门已同步到资产库。',
                )
                return redirect('requisition')
    return redirect('requisition')


@perm_required('manage_requisition')
def requisition_return(request, pk):
    record = get_object_or_404(Requisition, pk=pk)
    if request.method == 'POST':
        from django.db import transaction
        with transaction.atomic():
            record.status = '已归还'
            record.return_date = timezone.localdate()
            record.save()
        oplog.log(request, 'requisition', 'update', target=_asset_label(record.asset),
                  detail=(f'归还登记：领用人 {record.user}；归还日期 {record.return_date}；'
                          f'资产责任人/实际使用人已复位为「{Requisition.DEFAULT_OWNER}」'))
        messages.success(request, f'资产 {record.asset.name} 已归还。')
    return redirect('requisition')


@perm_required('manage_requisition')
def requisition_edit(request, pk):
    record = get_object_or_404(Requisition, pk=pk)
    if request.method == 'POST':
        before = {
            'user': record.user, 'actual_user': record.actual_user,
            'department': record.department.name if record.department_id else '',
            'purpose': record.purpose,
            'due_date': record.due_date.strftime('%Y-%m-%d') if record.due_date else '',
            'status': record.status,
        }
        record.user = request.POST.get('user', record.user).strip()
        record.actual_user = request.POST.get('actual_user', record.actual_user).strip()
        record.department_id = request.POST.get('department') or None
        record.purpose = request.POST.get('purpose', '').strip()
        record.due_date = request.POST.get('due_date') or None
        record.return_date = request.POST.get('return_date') or None
        # 状态归一化：存储态仅「借出/已归还」；逾期为派生状态，不允许手工入库
        raw_status = request.POST.get('status', '')
        record.status = '已归还' if raw_status == '已归还' else '借出'
        if not record.user:
            messages.error(request, '领用人不能为空。')
        elif not record.department_id:
            messages.error(request, '请选择领用部门。')
        elif not record.purpose:
            messages.error(request, '请填写领用用途。')
        elif not record.due_date:
            messages.error(request, '请选择预计归还日期。')
        else:
            record.save()
            after = {
                'user': record.user, 'actual_user': record.actual_user,
                'department': record.department.name if record.department_id else '',
                'purpose': record.purpose,
                'due_date': record.due_date.strftime('%Y-%m-%d') if record.due_date else '',
                'status': record.status,
            }
            diff_text = oplog.diff(before, after, {
                'user': '领用人', 'actual_user': '实际使用人', 'department': '领用部门',
                'purpose': '领用用途', 'due_date': '预计归还', 'status': '状态',
            })
            if diff_text:
                oplog.log(request, 'requisition', 'update',
                          target=_asset_label(record.asset), detail=diff_text)
            messages.success(request, '领用记录已更新。')
        return redirect('requisition')
    context = {
        'record': record,
        'departments': Department.objects.all(),
        'page_title': '编辑领用',
        'active': 'requisition',
    }
    return render(request, 'requisition_form.html', context)


@perm_required('manage_requisition')
def requisition_delete(request, pk):
    record = get_object_or_404(Requisition, pk=pk)
    if request.method == 'POST':
        from django.db import transaction
        with transaction.atomic():
            # 复位资产状态与归属的逻辑在 Requisition.delete() 中统一处理
            target = _asset_label(record.asset)
            detail = f'删除领用记录（领用人：{record.user}；状态：{record.display_status}）'
            record.delete()
        oplog.log(request, 'requisition', 'delete', target=target, detail=detail)
        messages.success(request, '领用记录已删除。')
    return redirect('requisition')


# ---------- 资产变更模块 ----------
@perm_required('view_change')
def change(request):
    ctype = request.GET.get('type', '')
    records = AssetChange.objects.select_related('asset').order_by('-change_date')
    if ctype:
        records = records.filter(change_type=ctype)

    paginator = Paginator(records, 12)
    page_obj = paginator.get_page(request.GET.get('page'))

    all_assets = Asset.objects.all()
    asset_options = [
        {'id': a.pk, 'asset_id': a.asset_id, 'name': a.name, 'status': a.status,
         'department': a.department.name if a.department else '',
         'responsible': a.responsible or '', 'user': a.user or ''}
        for a in all_assets
    ]

    context = {
        'records': page_obj.object_list,
        'page_obj': page_obj,
        'assets': all_assets,
        'asset_options': asset_options,
        'change_type_options': [{'id': t[0], 'text': t[0]} for t in AssetChange.CHANGE_TYPES],
        'departments': Department.objects.all(),
        'ctype': ctype,
        'page_title': '资产变更',
        'active': 'change',
    }
    return render(request, 'change.html', context)


@perm_required('view_change')
def change_export(request):
    """按当前筛选条件导出「资产变更」记录为 Excel。"""
    ctype = request.GET.get('type', '')
    records = AssetChange.objects.select_related('asset').order_by('-change_date')
    if ctype:
        records = records.filter(change_type=ctype)

    wb = Workbook()
    ws = wb.active
    ws.title = '资产变更'
    headers = ['资产编号', '资产名称', '变更类型', '变更原因', '变更前',
               '变更后', '责任人（变更前→后）', '实际使用人（变更前→后）',
               '经办人', '变更时间', '备注']
    ws.append(headers)

    def _pair(old, new):
        old, new = (old or '').strip(), (new or '').strip()
        if not old and not new:
            return ''
        return f'{old or "—"} → {new or "—"}'

    for r in records:
        ws.append([
            r.asset.asset_id,
            r.asset.name,
            r.change_type,
            r.reason,
            r.old_value,
            r.new_value,
            _pair(r.old_responsible, r.new_responsible),
            _pair(r.old_user, r.new_user),
            r.changed_by,
            timezone.localtime(r.change_date).strftime('%Y-%m-%d %H:%M') if r.change_date else '',
            r.note,
        ])
    for i, w in enumerate([14, 22, 12, 22, 16, 16, 20, 20, 12, 18, 22], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="change_export.xlsx"'
    wb.save(response)
    return response


@perm_required('manage_change')
def change_create(request):
    if request.method == 'POST':
        asset_id = request.POST.get('asset')
        change_type = request.POST.get('change_type', '部门转移')
        reason = request.POST.get('reason', '').strip()
        old_value = request.POST.get('old_value', '').strip()
        new_value = request.POST.get('new_value', '').strip()
        new_responsible = request.POST.get('new_responsible', '').strip()
        new_user = request.POST.get('new_user', '').strip()
        changed_by = request.POST.get('changed_by', '').strip()
        note = request.POST.get('note', '').strip()

        if not asset_id:
            messages.error(request, '请选择要变更的资产。')
        elif not reason:
            messages.error(request, '请填写变更原因。')
        elif not old_value:
            messages.error(request, '请填写变更前。')
        elif not new_value:
            messages.error(request, '请填写变更后。')
        elif not changed_by:
            messages.error(request, '请填写经办人。')
        else:
            asset = get_object_or_404(Asset, pk=asset_id)
            # 先校验可同步性，失败则不落库，避免产生从未生效的假变更记录
            ok, msg = _validate_change(change_type, new_value)
            if not ok:
                messages.error(request, msg)
            else:
                # 部门转移：自动记录变更前的责任人 / 实际使用人（历史快照）
                old_responsible = asset.responsible if change_type == '部门转移' else ''
                old_user = asset.user if change_type == '部门转移' else ''
                from django.db import transaction
                with transaction.atomic():
                    AssetChange.objects.create(
                        asset=asset, change_type=change_type, reason=reason,
                        old_value=old_value, new_value=new_value,
                        old_responsible=old_responsible, old_user=old_user,
                        new_responsible=new_responsible, new_user=new_user,
                        changed_by=changed_by, note=note,
                    )
                    ok2, msg2 = _apply_change_sync(
                        asset, change_type, new_value,
                        new_responsible=new_responsible, new_user=new_user,
                    )
                    if not ok2:
                        messages.warning(request, '变更记录已保存，但资产同步未完成：' + (msg2 or ''))
                    else:
                        messages.success(request, f'资产 {asset.name} 变更记录已保存，并同步到资产库。')
                    oplog.log(
                        request, 'change', 'create', target=_asset_label(asset),
                        detail=(f'{change_type}：{old_value} → {new_value}；原因：{reason}；'
                                f'经办人：{changed_by}'
                                + (f'；责任人：{old_responsible or "—"} → {new_responsible or "—"}'
                                   if change_type == '部门转移' else '')
                                + (f'；实际使用人：{old_user or "—"} → {new_user or "—"}'
                                   if change_type == '部门转移' else '')
                                + ('' if ok2 else '（资产同步未完成）')),
                    )
                return redirect('change')
    return redirect('change')


def _validate_change(change_type, new_value):
    """校验一条资产变更能否真正同步到台账。返回 (ok, error_msg)。

    - change_type 必须在系统支持的变更类型内（部门转移/状态变更）。
    - 部门转移：目标部门必须已存在于组织架构，禁止自动新建错误部门。
    - 状态变更：目标状态必须为资产状态枚举中的合法值。
    """
    new_value = (new_value or '').strip()
    valid_types = [c[0] for c in AssetChange.CHANGE_TYPES]
    if change_type not in valid_types:
        return False, f'变更类型「{change_type or "空"}」无效，仅支持：{"、".join(valid_types)}'
    if change_type == '部门转移':
        if not new_value:
            return False, '未填写「变更后」的目标部门。'
        if not Department.objects.filter(name=new_value).exists():
            return False, f'目标部门「{new_value}」在组织架构中不存在，请先在组织架构中创建或核对部门名称。'
    elif change_type == '状态变更':
        valid = [c[0] for c in Asset.STATUS_CHOICES]
        if new_value not in valid:
            return False, f'状态「{new_value or "空"}」无效，仅支持：{"、".join(valid)}'
    return True, None


def _apply_change_sync(asset, change_type, new_value, new_responsible=None, new_user=None):
    """按变更类型把变更应用到资产台账（应在 _validate_change 通过后调用）。返回 (ok, msg)。

    部门转移时可连带同步新责任人 / 新实际使用人（传 None 或空串表示不变）。
    """
    from django.db import transaction
    with transaction.atomic():
        if change_type == '部门转移':
            dept = Department.objects.get(name=new_value)
            asset.department = dept
            update_fields = ['department']
            if new_responsible:
                asset.responsible = new_responsible
                update_fields.append('responsible')
            if new_user:
                asset.user = new_user
                update_fields.append('user')
            asset.save(update_fields=update_fields)
        elif change_type == '状态变更':
            asset.status = new_value
            asset.save(update_fields=['status'])
        else:
            return False, '未知变更类型，未同步'
    return True, None


@perm_required('manage_change')
def change_update(request, pk):
    ch = get_object_or_404(AssetChange, pk=pk)
    if request.method == 'POST':
        reason = request.POST.get('reason', '').strip()
        ch.change_type = request.POST.get('change_type', ch.change_type)
        ch.old_value = request.POST.get('old_value', '').strip()
        ch.new_value = request.POST.get('new_value', '').strip()
        ch.old_responsible = request.POST.get('old_responsible', '').strip()
        ch.old_user = request.POST.get('old_user', '').strip()
        ch.new_responsible = request.POST.get('new_responsible', '').strip()
        ch.new_user = request.POST.get('new_user', '').strip()
        ch.changed_by = request.POST.get('changed_by', '').strip()
        ch.note = request.POST.get('note', '').strip()
        missing = []
        if not ch.change_type:
            missing.append('变更类型')
        if not reason:
            missing.append('变更原因')
        if not ch.old_value:
            missing.append('变更前')
        if not ch.new_value:
            missing.append('变更后')
        if not ch.changed_by:
            missing.append('经办人')
        if missing:
            messages.error(request, '请填写必填项：' + '、'.join(missing))
            return render(request, 'change_form.html', {'ch': ch, 'page_title': '编辑变更', 'active': 'change'})
        ch.reason = reason
        # 先校验可同步性，失败则不保存
        ok, msg = _validate_change(ch.change_type, ch.new_value)
        if not ok:
            messages.error(request, msg)
            return render(request, 'change_form.html', {'ch': ch, 'page_title': '编辑变更', 'active': 'change'})
        from django.db import transaction
        with transaction.atomic():
            ch.save()
            ok2, msg2 = _apply_change_sync(
                ch.asset, ch.change_type, ch.new_value,
                new_responsible=ch.new_responsible if ch.change_type == '部门转移' else None,
                new_user=ch.new_user if ch.change_type == '部门转移' else None,
            )
        if ok2:
            messages.success(request, '变更记录已更新，并同步到资产库。')
        else:
            messages.warning(request, '变更记录已更新，但资产同步未完成：' + (msg2 or ''))
        oplog.log(request, 'change', 'update', target=_asset_label(ch.asset),
                  detail=(f'修改变更记录 → {ch.change_type}：{ch.old_value} → {ch.new_value}；'
                          f'原因：{ch.reason}'
                          + ('' if ok2 else '（资产同步未完成）')))
        return redirect('change')
    context = {'ch': ch, 'page_title': '编辑变更', 'active': 'change'}
    return render(request, 'change_form.html', context)


@perm_required('manage_change')
def change_delete(request, pk):
    ch = get_object_or_404(AssetChange, pk=pk)
    if request.method == 'POST':
        target = _asset_label(ch.asset)
        detail = f'删除变更记录（{ch.change_type}：{ch.old_value} → {ch.new_value}）'
        ch.delete()
        oplog.log(request, 'change', 'delete', target=target, detail=detail)
        messages.success(request, '变更记录已删除。')
    return redirect('change')


# ---------- 用户管理模块（需 manage_users 权限） ----------
def _apply_role(user, role_obj):
    """给用户绑定角色，并同步 is_staff / is_superuser。"""
    from django.db import transaction
    with transaction.atomic():
        profile, _ = UserProfile.objects.get_or_create(user=user)
        profile.role = role_obj
        profile.save()
        user.is_superuser = bool(role_obj and role_obj.code == 'super_admin')
        user.is_staff = bool(role_obj and role_obj.can_access_admin)
        user.save(update_fields=['is_superuser', 'is_staff'])
    return profile


@perm_required('manage_users')
def user_list(request):
    users = User.objects.select_related('profile__role').order_by('id')
    context = {
        'users': users,
        'total_users': users.count(),
        'active_users': users.filter(is_active=True).count(),
        'admin_count': users.filter(is_staff=True).count(),
        'page_title': '用户管理',
        'active': 'users',
    }
    return render(request, 'user_list.html', context)


@perm_required('manage_users')
def user_create(request):
    roles = Role.objects.all()
    ctx = {
        'page_title': '新增用户', 'active': 'users', 'edit_user': None, 'roles': roles,
        'form_data': {},
    }
    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '')
        first_name = request.POST.get('first_name', '').strip()
        email = request.POST.get('email', '').strip()
        role_id = request.POST.get('role') or None

        # 失败时回显用户已填写内容
        ctx['form_data'] = {
            'username': username, 'password': password, 'first_name': first_name,
            'email': email, 'role_id': role_id,
        }

        pwd_err = _check_password_strength(None, password) if password else None

        if not username:
            messages.error(request, '用户名不能为空。')
        elif not password:
            messages.error(request, '密码不能为空。')
        elif not role_id:
            messages.error(request, '请为用户分配角色。')
        elif User.objects.filter(username=username).exists():
            messages.error(request, f'用户名 {username} 已存在。')
        elif pwd_err:
            messages.error(request, f'密码不符合要求：{pwd_err}')
        else:
            user = User.objects.create_user(
                username=username, password=password, email=email, first_name=first_name
            )
            role_obj = Role.objects.filter(pk=role_id).first()
            _apply_role(user, role_obj)
            oplog.log(request, 'user', 'create', target=f'用户 {username}',
                      detail=(f'姓名：{first_name or "（空）"}；邮箱：{email or "（空）"}；'
                              f'角色：{role_obj.name if role_obj else "（未分配）"}'))
            messages.success(request, f'用户 {username} 创建成功。')
            return redirect('user_list')
    return render(request, 'user_form.html', ctx)


@perm_required('manage_users')
def user_update(request, pk):
    user = get_object_or_404(User, pk=pk)
    roles = Role.objects.all()
    if request.method == 'POST':
        first_name = request.POST.get('first_name', '').strip()
        email = request.POST.get('email', '').strip()
        role_id = request.POST.get('role') or None
        is_active = request.POST.get('is_active') == '1'
        password = request.POST.get('password', '')
        role_obj = Role.objects.filter(pk=role_id).first() if role_id else None

        # 重置密码时同样套用强度校验
        if password:
            pwd_err = _check_password_strength(user, password)
            if pwd_err:
                messages.error(request, f'密码不符合要求：{pwd_err}')
                return redirect('user_list')

        # 内置账号保护：不可停用、不可降级（避免系统被锁死）
        target_profile = getattr(user, 'profile', None)
        if target_profile is not None and target_profile.is_builtin:
            is_active = True  # 内置账号始终启用
            if not (role_obj and role_obj.code == 'super_admin'):
                messages.error(request, '内置账号不可降级，必须保持「超级管理员」角色。')
                return redirect('user_list')

        # 自我保护：不能停用自己、不能移除自己的用户管理权限
        if user == request.user:
            if not is_active:
                messages.error(request, '不能停用当前登录的账号。')
            elif not (role_obj and role_obj.manage_users):
                messages.error(request, '不能移除自己的用户管理权限。')
            else:
                user.first_name = first_name
                user.email = email
                user.is_active = is_active
                _apply_role(user, role_obj)
                if password:
                    user.set_password(password)
                    user.save()
                oplog.log(request, 'user', 'update', target=f'用户 {user.username}',
                          detail=(f'（本人资料）姓名：{first_name or "（空）"}；邮箱：{email or "（空）"}；'
                                  f'角色：{role_obj.name if role_obj else "（未分配）"}'
                                  + ('；已重置密码' if password else '')))
                messages.success(request, '个人信息已更新。')
            return redirect('user_list')

        # 最后一个超管不能降级
        if user.is_superuser and not (role_obj and role_obj.code == 'super_admin'):
            if User.objects.filter(is_superuser=True).count() <= 1:
                messages.error(request, '系统至少需要保留一个超级管理员。')
                return redirect('user_list')

        user.first_name = first_name
        user.email = email
        user.is_active = is_active
        _apply_role(user, role_obj)
        if password:
            user.set_password(password)
            user.save()
        oplog.log(request, 'user', 'update', target=f'用户 {user.username}',
                  detail=(f'姓名：{first_name or "（空）"}；邮箱：{email or "（空）"}；'
                          f'启用：{"是" if is_active else "否"}；'
                          f'角色：{role_obj.name if role_obj else "（未分配）"}'
                          + ('；已重置密码' if password else '')))
        messages.success(request, f'用户 {user.username} 更新成功。')
        return redirect('user_list')

    context = {'page_title': '编辑用户', 'active': 'users', 'edit_user': user, 'roles': roles}
    return render(request, 'user_form.html', context)


@perm_required('manage_users')
def user_delete(request, pk):
    user = get_object_or_404(User, pk=pk)
    if request.method == 'POST':
        profile = getattr(user, 'profile', None)
        if profile is not None and profile.is_builtin:
            messages.error(request, '内置账号不可删除。')
        elif user == request.user:
            messages.error(request, '不能删除当前登录的账号。')
        elif user.is_superuser:
            messages.error(request, '不能删除超级管理员账号，请先将其降级。')
        else:
            uname = user.username
            user.delete()
            oplog.log(request, 'user', 'delete', target=f'用户 {uname}', detail='删除用户及其档案')
            messages.success(request, f'用户 {user.username} 已删除。')
    return redirect('user_list')


# ---------- 角色管理 ----------
_PERM_FIELDS = [
    ('can_access_admin', '访问后台'),
    ('view_assets', '查看资产库'), ('manage_assets', '管理资产库'),
    ('view_org', '查看组织架构'), ('manage_org', '管理组织架构'),
    ('view_requisition', '查看资产领用'), ('manage_requisition', '管理资产领用'),
    ('view_change', '查看资产变更'), ('manage_change', '管理资产变更'),
    ('manage_users', '管理用户/角色'),
]

_PERM_GROUPS = [
    {'title': '系统权限', 'items': [('can_access_admin', '访问后台管理')]},
    {'title': '资产库', 'items': [('view_assets', '查看'), ('manage_assets', '管理')]},
    {'title': '组织架构', 'items': [('view_org', '查看'), ('manage_org', '管理')]},
    {'title': '资产领用', 'items': [('view_requisition', '查看'), ('manage_requisition', '管理')]},
    {'title': '资产变更', 'items': [('view_change', '查看'), ('manage_change', '管理')]},
    {'title': '用户/角色', 'items': [('manage_users', '管理用户/角色')]},
]


@perm_required('manage_users')
def role_list(request):
    roles = Role.objects.prefetch_related('profiles').order_by('id')
    context = {
        'roles': roles,
        'perm_fields': _PERM_FIELDS,
        'page_title': '角色管理',
        'active': 'roles',
    }
    return render(request, 'role_list.html', context)


@perm_required('manage_users')
def role_create(request):
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        code = request.POST.get('code', '').strip()
        description = request.POST.get('description', '').strip()
        if not name or not code:
            messages.error(request, '角色名称与编码不能为空。')
        elif Role.objects.filter(code=code).exists():
            messages.error(request, f'编码 {code} 已存在。')
        else:
            role = Role.objects.create(name=name, code=code, description=description)
            _save_perm_fields(role, request.POST)
            oplog.log(request, 'user', 'create', target=f'角色 {name}',
                      detail=f'编码：{code}；权限：{"、".join(_perm_labels(role)) or "（无）"}')
            messages.success(request, f'角色 {name} 创建成功。')
            return redirect('role_list')
    context = {'page_title': '新增角色', 'active': 'roles', 'edit_role': None,
               'perm_fields': _PERM_FIELDS, 'perm_groups': _PERM_GROUPS}
    return render(request, 'role_form.html', context)


@perm_required('manage_users')
def role_update(request, pk):
    role = get_object_or_404(Role, pk=pk)
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        description = request.POST.get('description', '').strip()
        if not name:
            messages.error(request, '角色名称不能为空。')
        else:
            role.name = name
            role.description = description
            role.save()
            _save_perm_fields(role, request.POST)
            oplog.log(request, 'user', 'update', target=f'角色 {name}',
                      detail=f'权限：{"、".join(_perm_labels(role)) or "（无）"}')
            messages.success(request, f'角色 {name} 更新成功。')
            return redirect('role_list')
    context = {'page_title': '编辑角色', 'active': 'roles', 'edit_role': role,
               'perm_fields': _PERM_FIELDS, 'perm_groups': _PERM_GROUPS}
    return render(request, 'role_form.html', context)


def _save_perm_fields(role, post):
    for field, _label in _PERM_FIELDS:
        setattr(role, field, post.get(field) == '1')
    role.save()


@perm_required('manage_users')
def role_delete(request, pk):
    role = get_object_or_404(Role, pk=pk)
    if request.method == 'POST':
        if role.is_system:
            messages.error(request, '系统内置角色不能删除。')
        elif role.profiles.exists():
            messages.error(request, f'还有 {role.profiles.count()} 个用户使用该角色，请先调整这些用户的角色。')
        else:
            rname = role.name
            role.delete()
            oplog.log(request, 'user', 'delete', target=f'角色 {rname}', detail='删除角色')
            messages.success(request, f'角色 {role.name} 已删除。')
    return redirect('role_list')


def _perm_labels(role):
    """把角色已开启的权限翻译成中文标签列表，用于操作日志。"""
    labels = []
    for field, label in _PERM_FIELDS:
        if getattr(role, field, False):
            labels.append(label)
    return labels


# ---------- 资产二维码 / 标签打印 / 扫码信息卡 ----------
QR_CACHE_SECONDS = 24 * 3600


def _cached_qr_svg(payload):
    """按内容缓存二维码图片。

    二维码内容里带了资产快照，所以"内容相同"就等于"图相同"——
    直接用内容做缓存键：资产没变就命中缓存，资产一改内容就变、自然重算，
    不需要任何失效逻辑。生成一张要 ~30ms，标签页动辄上百张，这层缓存很关键。
    """
    key = 'assetqr:' + hashlib.md5(payload.encode('utf-8')).hexdigest()
    svg = cache.get(key)
    if svg is None:
        svg = qr_svg(payload)
        cache.set(key, svg, QR_CACHE_SECONDS)
    return svg


@perm_required('view_assets')
def asset_qr(request, pk):
    """返回单个资产的二维码图片（SVG）。

    默认把「资产信息快照」编码进二维码：扫码后直接看到一个只读信息卡
    （``asset_card``），**不需要登录，也不会落到资产库列表页**。
    信息跟着二维码走，所以标签贴出去之后，没账号的人也能看清这是什么资产。

    ``?data=code`` 时退化为只编码资产编号纯文本（兼容扫码枪/旧标签）。

    二维码在本机离线生成，不依赖任何外部服务。响应带 ETag 且要求浏览器
    每次校验：资产一改，扫码内容立刻变新，不会把旧快照打进标签。
    """
    asset = get_object_or_404(Asset.objects.select_related('department'), pk=pk)
    if request.GET.get('data') == 'code':
        payload = asset.asset_id
    else:
        payload = snapshot.card_payload(request.build_absolute_uri(reverse('asset_card')), asset)

    response = HttpResponse(_cached_qr_svg(payload), content_type='image/svg+xml')
    response['ETag'] = '"%s"' % hashlib.md5(payload.encode('utf-8')).hexdigest()
    response['Cache-Control'] = 'private, no-cache, max-age=0'
    return get_conditional_response(request, etag=response['ETag'], response=response) or response


@require_GET
def asset_card(request):
    """扫码落地页：直接显示二维码里携带的资产信息快照。

    * **免登录**：扫码的人不需要系统账号；
    * **只读**：没有任何编辑/跳转入口，只有一个返回说明；
    * **不查库**：页面内容全部来自 ``?d=`` 参数，改参数也枚举不出别的资产。

    参数缺失或损坏时给出可读的提示页，而不是报错栈。
    """
    data = snapshot.decode(request.GET.get('d', ''))
    fields = []
    if data:
        for key in snapshot.CARD_FIELDS:
            value = (data.get(key) or '').strip()
            if not value:
                continue
            if key == 'price':
                # 台账里 0 元多半是"未估值"，显示 ¥0.00 反而像出错，直接略过。
                if value in ('0.00', '0'):
                    continue
                value = f'¥ {value}'
            fields.append({'key': key, 'label': snapshot.LABELS[key], 'value': value})

    context = {
        'a': data,
        'fields': fields,
        'tone': snapshot.status_tone(data['status']) if data else 'muted',
    }
    response = render(request, 'asset_card.html', context)
    response['X-Robots-Tag'] = 'noindex, nofollow'
    return response


@perm_required('view_assets')
def asset_labels(request):
    """标签打印页：把所选（或当前筛选）资产排成便于裁剪张贴的标签。

    取值优先级：
    1. ``?ids=1,2,3`` —— 资产库勾选后点「打印标签」（最常用）；
    2. 否则复用资产库的筛选参数（状态/类别/部门/关键字），最多 200 张。
    """
    ids_raw = (request.GET.get('ids') or '').strip()
    if ids_raw:
        ids = [int(x) for x in ids_raw.split(',') if x.strip().isdigit()]
        assets = list(
            Asset.objects.select_related('department').filter(pk__in=ids).order_by('asset_id')
        )
    else:
        qs, _ = _filtered_assets(request)
        assets = list(qs.order_by('asset_id')[:200])

    context = {
        'assets': assets,
        'page_title': '资产标签打印',
        'active': 'library',
    }
    return render(request, 'asset_labels.html', context)


# ---------- 资产附件 ----------
@perm_required('manage_assets')
def attachment_upload(request, pk):
    """为资产上传附件（支持一次多选）。"""
    asset = get_object_or_404(Asset, pk=pk)
    if request.method == 'POST':
        files = request.FILES.getlist('files')
        if not files:
            messages.error(request, '请先选择要上传的文件。')
            return redirect('asset_update', pk=pk)

        allowed_text = '、'.join(AssetAttachment.ALLOWED_EXT)
        max_mb = AssetAttachment.MAX_SIZE // (1024 * 1024)
        ok, errors = 0, []
        for f in files:
            ext = os.path.splitext(f.name)[1].lower()
            if ext not in AssetAttachment.ALLOWED_EXT:
                errors.append(f'「{f.name}」格式不允许（支持：{allowed_text}）')
                continue
            if f.size > AssetAttachment.MAX_SIZE:
                errors.append(f'「{f.name}」超过 {max_mb}MB 上限')
                continue
            att = AssetAttachment.objects.create(
                asset=asset, file=f, name=f.name[:200], size=f.size,
                uploaded_by=(request.user.first_name or request.user.username),
            )
            oplog.log(request, 'attachment', 'create', target=_asset_label(asset),
                      detail=f'上传附件：{att.name}（{att.size_h}）')
            ok += 1

        if ok:
            messages.success(request, f'已上传 {ok} 个附件。')
        for e in errors:
            messages.error(request, e)
    return redirect('asset_update', pk=pk)


@perm_required('manage_assets')
def attachment_delete(request, pk):
    att = get_object_or_404(AssetAttachment, pk=pk)
    asset_pk = att.asset_id
    if request.method == 'POST':
        name = att.name
        asset_label = _asset_label(att.asset)
        try:
            att.file.delete(save=False)  # 一并删除物理文件
        except Exception:  # noqa: BLE001 — 文件已丢失也要能删记录
            pass
        att.delete()
        oplog.log(request, 'attachment', 'delete', target=asset_label, detail=f'删除附件：{name}')
        messages.success(request, f'附件「{name}」已删除。')
    return redirect('asset_update', pk=asset_pk)


@perm_required('view_assets')
def attachment_download(request, pk):
    """附件下载：走登录 + 权限校验，不直接暴露媒体目录。"""
    att = get_object_or_404(AssetAttachment, pk=pk)
    try:
        path = att.file.path
    except Exception:  # noqa: BLE001
        raise Http404('附件路径无效')
    if not os.path.exists(path):
        raise Http404('附件文件已丢失')
    return FileResponse(
        open(path, 'rb'), as_attachment=True,
        filename=att.name or os.path.basename(path),
    )


# ---------- 操作日志 ----------
@admin_required
def oplog_list(request):
    """操作日志查询页（管理员）。支持按模块/动作/关键字/日期区间筛选。"""
    category = request.GET.get('category', '')
    action = request.GET.get('action', '')
    keyword = request.GET.get('q', '').strip()
    start = request.GET.get('start', '').strip()
    end = request.GET.get('end', '').strip()

    qs = OperationLog.objects.all()
    if category:
        qs = qs.filter(category=category)
    if action:
        qs = qs.filter(action=action)
    if keyword:
        qs = qs.filter(
            Q(operator__icontains=keyword) | Q(target__icontains=keyword)
            | Q(detail__icontains=keyword) | Q(ip__icontains=keyword)
        )
    if start:
        qs = qs.filter(created_at__date__gte=start)
    if end:
        qs = qs.filter(created_at__date__lte=end)

    paginator = Paginator(qs, 50)
    page_obj = paginator.get_page(request.GET.get('page'))

    context = {
        'page_obj': page_obj,
        'records': page_obj.object_list,
        'category_options': OperationLog.CATEGORY_CHOICES,
        'action_options': OperationLog.ACTION_CHOICES,
        'filters': {'category': category, 'action': action, 'q': keyword, 'start': start, 'end': end},
        'total_logs': OperationLog.objects.count(),
        'page_title': '操作日志',
        'active': 'oplog',
    }
    return render(request, 'oplog.html', context)
