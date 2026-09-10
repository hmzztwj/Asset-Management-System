"""把「迈德固定资产在线表.xlsx」原始台账转换为「资产库 - 导入」模板格式。

原始表是一张 30 列的在线导出表（表头在第 3 行），字段命名和页面导入模板不一致，
直接上传无法识别。本脚本负责把它「翻译」成页面导入要求的 13 列模板，方便在
「资产库 → 导入资产」直接上传。

用法:
    python convert_to_template.py                     # 用默认路径转换
    python convert_to_template.py --fill-user         # 实际使用人为空/「无」时用责任人填充
    python convert_to_template.py --create-departments # 顺便把缺失的部门写入数据库
    python convert_to_template.py --src 某表.xlsx --out 某输出.xlsx

列映射（原始列 → 模板列）:
    财务编码[20]      → 资产编号（原始「资产编号」列常为空，以财务编码为准）
    资产分类[1]       → 资产名称 + 资产类别（类别按别名表归一化）
    责任部门[4]       → 所属部门
    责任人员[5]       → 责任人
    实际使用人[7]     → 实际使用人（「无」视为空）
    存放位置[8]       → 存放位置
    资产状态[0]       → 使用状态（按别名表归一化）
    序列号[11]        → 序列号
    笔记本尺寸/显示器 → 规格
    CPU/内存/硬盘/显卡/尺寸/显示器 → 配置
    资产原值[21]      → 价格
"""
import argparse
import os
import sys

import django

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "assets_system.settings")
django.setup()

from openpyxl import Workbook, load_workbook  # noqa: E402

from assets.models import Asset, Department  # noqa: E402

DEFAULT_SRC = os.path.join("实际数据", "迈德固定资产在线表.xlsx")
DEFAULT_OUT = os.path.join("实际数据", "导入模板_资产数据.xlsx")

HEADERS = [
    "资产编号", "资产名称", "资产类别", "所属部门", "责任人", "实际使用人",
    "存放位置", "使用状态", "序列号", "规格", "配置", "价格", "购置日期",
]

# 原始列下标
C_STATUS, C_CATEGORY, C_COMPANY, C_DEPT = 0, 1, 3, 4
C_RESPONSIBLE, C_USER, C_LOCATION, C_SERIAL = 5, 7, 8, 11
C_CPU, C_RAM, C_DISK, C_GPU, C_SIZE, C_MONITOR = 14, 15, 16, 17, 18, 19
C_FIN_CODE, C_PRICE = 20, 21

STATUS_MAP = {**{c[0]: c[0] for c in Asset.STATUS_CHOICES}, **Asset.STATUS_ALIAS}
CAT_MAP = {**{c[0]: c[0] for c in Asset.CATEGORY_CHOICES}, **Asset.CATEGORY_ALIAS}


def cell(row, idx):
    if idx >= len(row):
        return ''
    v = row[idx]
    if v is None:
        return ''
    return str(v).strip()


def clean(value):
    """把「无」「/ 」等占位值统一视为空。"""
    text = (value or '').strip()
    return '' if text in ('无', '/', '—', '-') else text


def build_config(row):
    parts = [clean(cell(row, i)) for i in
             (C_CPU, C_RAM, C_DISK, C_GPU, C_SIZE, C_MONITOR)]
    return ' / '.join([p for p in parts if p])


def build_spec(row):
    return clean(cell(row, C_SIZE)) or clean(cell(row, C_MONITOR))


def convert(src, out, fill_user=False, create_departments=False):
    wb_in = load_workbook(src, data_only=True, read_only=True)
    ws_in = wb_in.active
    rows = list(ws_in.iter_rows(values_only=True))
    data = rows[3:]  # 表头在第 3 行

    records = []
    skipped = []          # 必填项缺失
    normalized = {'category': 0, 'status': 0}
    user_filled = 0
    seen_codes = set()
    duplicates = []

    for row in data:
        if not any(c is not None and str(c).strip() for c in row):
            continue

        asset_id = cell(row, C_FIN_CODE)
        if not asset_id:
            continue
        if asset_id in seen_codes:
            duplicates.append(asset_id)
        seen_codes.add(asset_id)

        raw_cat = cell(row, C_CATEGORY)
        name = raw_cat or cell(row, 2) or '资产'
        category = CAT_MAP.get(raw_cat, '其他')
        if category != raw_cat:
            normalized['category'] += 1

        raw_status = cell(row, C_STATUS)
        status = STATUS_MAP.get(raw_status, '库存')
        if status != raw_status:
            normalized['status'] += 1

        dept_name = cell(row, C_DEPT)
        responsible = cell(row, C_RESPONSIBLE)
        user = clean(cell(row, C_USER))
        location = clean(cell(row, C_LOCATION))

        if fill_user and not user:
            user = responsible
            user_filled += 1

        missing = [label for label, val in
                   (('所属部门', dept_name), ('责任人', responsible),
                    ('实际使用人', user), ('存放位置', location)) if not val]
        if missing:
            skipped.append((asset_id, '、'.join(missing)))
            continue

        price = 0
        try:
            price = float(cell(row, C_PRICE)) if cell(row, C_PRICE) else 0
        except ValueError:
            price = 0

        records.append([
            asset_id, name, category, dept_name, responsible, user,
            location, status, cell(row, C_SERIAL), build_spec(row),
            build_config(row), price, '',
        ])

    # 部门存在性检查（页面导入要求部门已存在，否则该行会被跳过）
    existing = set(Department.objects.values_list('name', flat=True))
    used_depts = sorted({r[3] for r in records})
    missing_depts = [d for d in used_depts if d not in existing]

    if create_departments and missing_depts:
        company_name = cell(data[0], C_COMPANY) or '公司'
        company, _ = Department.objects.get_or_create(
            name=company_name, defaults={'description': '责任公司'})
        for dname in missing_depts:
            Department.objects.get_or_create(name=dname, defaults={'parent': company})
        existing = set(Department.objects.values_list('name', flat=True))
        missing_depts = [d for d in used_depts if d not in existing]

    # 写出模板
    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = '资产导入模板'
    ws_out.append(HEADERS)
    for rec in records:
        ws_out.append(rec)
    widths = [18, 22, 12, 16, 12, 14, 18, 10, 16, 14, 34, 12, 14]
    for i, w in enumerate(widths, start=1):
        ws_out.column_dimensions[ws_out.cell(row=1, column=i).column_letter].width = w
    ws_out.freeze_panes = 'A2'
    wb_out.save(out)

    # ---------- 报告 ----------
    print('=' * 64)
    print(f'源文件：{src}')
    print(f'输出文件：{out}')
    print('-' * 64)
    print(f'原始数据行数    : {len(data)}')
    print(f'成功转换        : {len(records)}')
    print(f'类别归一化      : {normalized["category"]} 条')
    print(f'状态归一化      : {normalized["status"]} 条')
    if user_filled:
        print(f'用责任人补使用人: {user_filled} 条')
    print(f'必填项缺失跳过  : {len(skipped)} 条')
    for aid, why in skipped[:20]:
        print(f'    - {aid}：缺 {why}')
    if len(skipped) > 20:
        print(f'    ...（其余 {len(skipped) - 20} 条省略）')

    if duplicates:
        print(f'重复资产编号    : {len(duplicates)} 个 -> {duplicates[:10]}')

    if missing_depts:
        print(f'数据库中不存在的部门（这些行导入时会被跳过）: {missing_depts}')
        print('    → 可先在「组织架构」里建好，或加 --create-departments 自动创建')
    else:
        print('部门检查        : 全部存在于数据库，可直接导入')

    print('-' * 64)
    print('下一步：打开系统「资产库 → 导入资产」，上传上面这个输出文件即可。')
    print('=' * 64)

    return len(records), len(skipped), missing_depts


def main():
    parser = argparse.ArgumentParser(description='原始台账 → 页面导入模板')
    parser.add_argument('--src', default=DEFAULT_SRC, help='原始 Excel 路径')
    parser.add_argument('--out', default=DEFAULT_OUT, help='输出 Excel 路径')
    parser.add_argument('--fill-user', action='store_true',
                        help='实际使用人为空/「无」时用责任人填充')
    parser.add_argument('--create-departments', action='store_true',
                        help='把缺失的部门自动写入数据库（父级为责任公司）')
    args = parser.parse_args()

    if not os.path.exists(args.src):
        print(f'找不到源文件：{args.src}')
        sys.exit(1)
    convert(args.src, args.out, fill_user=args.fill_user,
            create_departments=args.create_departments)


if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    main()
